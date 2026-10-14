"""Build the updated project report from the reviewed report and latest metrics."""

from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "reports" / "Итоговый_отчет_SBER.docx"
OUTPUT = ROOT / "reports" / "Итоговый_отчет_SBER_обновленный.docx"


def replace_paragraph_containing(document: Document, needle: str, replacement: str) -> None:
    for paragraph in document.paragraphs:
        if needle in paragraph.text:
            paragraph.text = replacement
            return
    raise ValueError(f"Report paragraph not found: {needle}")


def main() -> None:
    document = Document(SOURCE)
    replace_paragraph_containing(
        document,
        "В качестве второго, слабоструктурированного источника подключена",
        "В качестве второго, слабоструктурированного источника подключена RSS-лента пресс-релизов Банка России. Она содержит лишь 15 публикаций за 11 сентября — 8 октября 2026 года, покрывает 0,50% торговых сессий и не является новостным архивом SBER; поэтому текстовые признаки не включались в обучение.",
    )
    replace_paragraph_containing(
        document,
        "Сравнение на хронологических выборках включало",
        "Сравнение включало наивный прогноз и три модели машинного обучения. Для каждого ML-класса набор признаков выбирался по MAE на валидации; выиграла линейная регрессия с пятью ценовыми признаками, MAE 3,29 руб. На отложенном тесте её MAE составила 2,43 руб., MAPE — 0,82%. Результат не доказывает возможность получения торговой прибыли.",
    )
    replace_paragraph_containing(
        document,
        "Признаки включают цену и объём текущей сессии",
        "Исходный набор включает цену и объём текущей сессии, доходности и лаги на 1, 2 и 5 сессий, средние цены за 5 и 20 сессий, волатильность и календарные признаки. День недели, месяц и день месяца представлены парами sin/cos, что учитывает близость соседних точек календарного цикла. Цель — закрытие следующей сессии.",
    )
    replace_paragraph_containing(
        document,
        "На валидации базовый прогноз оказался точнее всех проверенных ML-моделей",
        "По валидационной MAE выбрана линейная регрессия с набором price_levels_only: цена текущей сессии, лаги цены на 1 и 5 сессий и средние цены за 5 и 20 сессий. Её MAE равна 3,29 руб. против 3,31 руб. у наивного прогноза. Разница невелика; выбранная конфигурация проверена один раз на тесте, который не участвовал в отборе. Тестовые метрики: MAE 2,43 руб., RMSE 3,44 руб., MAPE 0,82%, R² 0,959. Метрики относятся к цене, а не к доходности портфеля.",
    )
    replace_paragraph_containing(
        document,
        "Последняя 20-сессионная MAE составила",
        "Последняя 20-сессионная MAE составила 2,58 руб. против 3,29 руб. на валидации; отношение 0,79, порог проверки 1,5 не достигнут. Ошибка выросла в июне–августе 2026 года. Порог — сигнал аналитику, а не доказательство дрейфа.",
    )
    replace_paragraph_containing(
        document,
        "Streamlit-дашборд показывает последнюю котировку",
        "Streamlit-дашборд показывает последнюю котировку, прогноз выбранной модели на следующую сессию, историю, тестовые прогнозы, ошибки, покрытие RSS, сравнение групп признаков и локальное объяснение текущего прогноза.",
    )
    replace_paragraph_containing(
        document,
        "Сравнение и тестовые прогнозы сохраняются",
        "Сравнение и тестовые прогнозы сохраняются в reports/model_comparison.json и reports/test_predictions.csv. reports/feature_analysis.json содержит валидационную абляцию групп признаков, две оценки важности деревьев, их сопоставление между validation и тестовым периодом и кривые частичной зависимости. Сводка мониторинга записывается в reports/monitoring.json; признаки и проверки — в data/processed/features.csv, data/processed/latest_features.json и data/processed/quality_report.json.",
    )
    # Renumber the final sections to make room for feature analysis.
    replace_paragraph_containing(document, "7 Интерпретация и мониторинг", "8 Интерпретация и мониторинг")
    replace_paragraph_containing(document, "8 Дашборд и воспроизведение", "9 Дашборд и воспроизведение")
    replace_paragraph_containing(document, "9 Ограничения и дальнейшая работа", "10 Ограничения и дальнейшая работа")
    replace_paragraph_containing(
        document,
        "Наивная модель переносит последнее закрытие",
        "Локальное объяснение текущего прогноза строится заменой по одному признаку на его медиану обучающей таблицы при фиксированных остальных признаках. Для линейной регрессии эти отклонения дают точное аддитивное разложение; для нелинейных моделей это диагностические изменения, которые не обязаны складываться.",
    )
    replace_paragraph_containing(
        document,
        "RSS покрывает лишь последние недели",
        "RSS Банка России покрывает лишь последние недели и не является новостной лентой SBER; историческое влияние текста на прогноз проверить нельзя.",
    )
    replace_paragraph_containing(
        document,
        "Создан воспроизводимый прототип сбора и анализа данных SBER.",
        "Создан воспроизводимый прототип сбора и анализа данных SBER. Отбор групп признаков на валидации выбрал линейную регрессию с пятью ценовыми признаками; тестовая MAE составила 2,43 руб. RSS-истории недостаточно для оценки влияния новостей. Для следующего этапа потребуется автоматизированный архив публикаций и мониторинг распределения входных данных в эксплуатации.",
    )

    for table in document.tables:
        for row in table.rows:
            if row.cells[0].text.strip() == "Линейная регрессия":
                row.cells[2].text = "Линейная модель, набор price_levels_only выбран по validation MAE"
            if row.cells[0].text.strip() == "Случайный лес":
                row.cells[2].text = "Ансамбль деревьев; набор признаков выбран по validation MAE"
            if row.cells[0].text.strip() == "Градиентный бустинг":
                row.cells[2].text = "Последовательный ансамбль; набор признаков выбран по validation MAE"
            if row.cells[0].text.strip() == "Линейная регрессия" and len(row.cells) == 4:
                row.cells[1].text = "3,29"
                row.cells[2].text = "4,68"
                row.cells[3].text = "1,15%"
            elif row.cells[0].text.strip() == "Случайный лес" and len(row.cells) == 4:
                row.cells[1].text = "3,39"
                row.cells[2].text = "4,82"
                row.cells[3].text = "1,19%"
            elif row.cells[0].text.strip() == "Градиентный бустинг" and len(row.cells) == 4:
                row.cells[1].text = "3,39"
                row.cells[2].text = "4,78"
                row.cells[3].text = "1,19%"

    marker = next((p for p in document.paragraphs if p.text == "8 Интерпретация и мониторинг"), None)
    if marker is None:
        raise ValueError("Could not locate feature-analysis insertion point")
    heading = marker.insert_paragraph_before("7 Отбор признаков и интерпретация", style="Heading 1")
    heading.paragraph_format.keep_with_next = True
    paragraphs = [
        "Для каждого класса ML-модели проверены пять наборов признаков. Выбор делался по MAE на validation: линейная регрессия выбрала price_levels_only и достигла MAE 3,29 руб.; случайный лес — price_levels_only, 3,39 руб.; градиентный бустинг — price_levels_only, 3,39 руб. Наивный прогноз показал 3,31 руб. и не зависит от этого отбора. Разбиение сохранено по времени, строки с целевой датой на границах периодов удалены из обучения.",
        "Встроенная важность деревьев сопоставлена с permutation importance, которая измеряет рост MAE после перемешивания признака. Permutation importance рассчитана отдельно на validation и тесте на моделях, обученных только на train. Тестовая диагностика выполняется после фиксации итоговой модели и не влияет на отбор.",
        "Для цены и дневной доходности рассчитаны одномерные кривые частичной зависимости на распределении validation. Они показывают усреднённое поведение модели, а не причинное влияние. Текущий прогноз выбранной линейной модели объясняется локально через отклонение каждого признака от медианы обучающего набора.",
        "Итоги расчётов сохранены в reports/feature_analysis.json и показаны в дашборде.",
    ]
    for text in paragraphs:
        p = marker.insert_paragraph_before(text, style="Normal")
        p.paragraph_format.keep_together = True

    # Add the investigated public issuer-disclosure page as a source, clearly marked as not used in training.
    document.add_paragraph(
        "[5] ПРАЙМ. Публичная страница раскрытия информации ПАО Сбербанк (источник изучен, архив не подключён к обучению). https://disclosure.1prime.ru/Portal/Default.aspx?emid=7707083893"
    )
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
