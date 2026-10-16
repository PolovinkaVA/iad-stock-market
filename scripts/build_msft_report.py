"""Build the reader-facing MSFT report with a title page from a supplied DOCX.

Uses python-docx from the document workspace runtime. The reference is read only.
"""
import argparse
import copy
import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

ROOT=Path(__file__).resolve().parents[1]

def black(run,size=None,bold=None):
    run.font.name='Times New Roman';run.font.color.rgb=RGBColor(0,0,0)
    if size is not None:run.font.size=Pt(size)
    if bold is not None:run.bold=bold

def replace(p,text,bold=False):
    p.clear();black(p.add_run(text),14,bold)

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--reference',type=Path,required=True)
    parser.add_argument('--output',type=Path,default=ROOT/'reports/Отчет_ИАД_MSFT_Половинка_ВА.docx')
    args=parser.parse_args()
    initial_hash=hashlib.sha256(args.reference.read_bytes()).hexdigest()
    doc=Document(args.reference)
    body=doc._element.body;keep=list(body)[:15]
    for child in list(body):
        if child not in keep and child.tag!=qn('w:sectPr'):body.remove(child)
    # The reference is used only for its cover. Drop relationships to lab figures.
    for rid,rel in list(doc.part.rels.items()):
        if rel.reltype.endswith('/image'):doc.part.drop_rel(rid)
    for p in doc.paragraphs[:15]:
        p.paragraph_format.first_line_indent=Cm(0)
        if p.paragraph_format.line_spacing is None:p.paragraph_format.line_spacing=1.0
        if p.paragraph_format.space_after is None:p.paragraph_format.space_after=Pt(0)
    replace(doc.paragraphs[8],'ОТЧЕТ ПО СЕМЕСТРОВОЙ РАБОТЕ',True)
    replace(doc.paragraphs[9],'по дисциплине «Интеллектуальный анализ данных»')
    replace(doc.paragraphs[10],'на тему «Прогнозирование цены акций Microsoft')
    replace(doc.paragraphs[11],'на основе котировок и корпоративных публикаций»')
    replace(doc.paragraphs[13],'Проверил:\nк. т. н., доцент\nТетерин Н. Н.')
    doc.paragraphs[14].paragraph_format.space_before=Pt(75)
    for name in ('Normal','Title','Heading 1','Heading 2','Heading 3','Caption'):
        s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(14);s.font.color.rgb=RGBColor(0,0,0)
        fonts=s.element.find('.//'+qn('w:rFonts'))
        if fonts is not None:
            for attr in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
                fonts.attrib.pop(qn('w:'+attr),None)
        s.paragraph_format.first_line_indent=Cm(0)
        borders=s.element.find('.//'+qn('w:pBdr'))
        if borders is not None:borders.getparent().remove(borders)
    normal=doc.styles['Normal'];normal.paragraph_format.line_spacing=1.15
    normal.paragraph_format.space_after=Pt(0);normal.paragraph_format.first_line_indent=Cm(1.25)
    normal.paragraph_format.widow_control=True
    for level in (1,2,3):
        s=doc.styles[f'Heading {level}'];s.font.bold=True
        s.paragraph_format.space_before=Pt(12 if level>1 else 0);s.paragraph_format.space_after=Pt(8)
        s.paragraph_format.keep_with_next=True;s.paragraph_format.line_spacing=1.15
        s.paragraph_format.page_break_before=level==1
        s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    doc.styles['Title'].font.bold=True
    doc.paragraphs[8].style=doc.styles['Title']
    caption_style=doc.styles['Caption'];caption_style.font.size=Pt(11);caption_style.font.italic=False
    caption_style.paragraph_format.line_spacing=1.0;caption_style.paragraph_format.space_after=Pt(8)
    # A short one-level TOC keeps the structure navigable without consuming the report.
    toc_title=doc.add_paragraph('СОДЕРЖАНИЕ')
    toc_title.paragraph_format.page_break_before=True;toc_title.paragraph_format.first_line_indent=Cm(0)
    toc_title.paragraph_format.space_after=Pt(12)
    for run in toc_title.runs:black(run,14,True)
    p=doc.add_paragraph();p.paragraph_format.first_line_indent=Cm(0)
    r=p.add_run();begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');r._r.append(begin)
    instr=OxmlElement('w:instrText');instr.set(qn('xml:space'),'preserve');instr.text=' TOC \\o "1-1" \\h \\z '
    r._r.append(instr);sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');r._r.append(sep)
    r.add_text('Содержание обновляется при открытии документа')
    end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');r._r.append(end)
    assets=json.loads((ROOT/'reports/report_tables.json').read_text(encoding='utf-8'))
    table_no=figure_no=0

    def add_table(item):
        nonlocal table_no
        table_no+=1
        p=doc.add_paragraph(f'Таблица {table_no} — {item["caption"]}',style='Caption');p.paragraph_format.keep_with_next=True
        n=len(item['header']);tab=doc.add_table(rows=1,cols=n);tab.alignment=WD_TABLE_ALIGNMENT.CENTER;tab.autofit=False
        total=16.5;widths=([5.2,11.3] if n==2 else ([4.6,5.95,5.95] if n==3 else ([6.3]+[(total-6.3)/(n-1)]*(n-1))))
        for col,width in zip(tab.columns,widths):col.width=Cm(width)
        for row in [item['header']]+item['rows']:
            cells=tab.rows[0].cells if row is item['header'] else tab.add_row().cells
            for j,(cell,value) in enumerate(zip(cells,row)):
                cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1.0
                p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2)
                if row is item['header'] or row is item['rows'][0]:p.paragraph_format.keep_with_next=True
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                black(p.add_run(str(value)),11,row is item['header'])
                pr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
                for edge in ('top','left','bottom','right'):
                    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
                pr.append(borders)
                margins=OxmlElement('w:tcMar')
                for edge in ('top','left','bottom','right'):
                    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'60');e.set(qn('w:type'),'dxa');margins.append(e)
                pr.append(margins)
                if row is item['header']:
                    fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E8EDF2');pr.append(fill)
            trpr=cells[0]._tc.getparent().get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
        head=OxmlElement('w:tblHeader');tab.rows[0]._tr.get_or_add_trPr().append(head)
        after=doc.add_paragraph();after.paragraph_format.space_after=Pt(2);after.paragraph_format.line_spacing=0.4

    def add_figure(key):
        nonlocal figure_no
        figure_no+=1
        p=doc.add_paragraph();p.paragraph_format.first_line_indent=Cm(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(6);p.paragraph_format.line_spacing=1.0
        run=p.add_run();pic=run.add_picture(str(ROOT/'reports/figures'/f'{key}.png'),width=Cm(16.0))
        pic._inline.docPr.set('descr',assets['figures'][key])
        caption=doc.add_paragraph(f'Рисунок {figure_no} — {assets["figures"][key]}',style='Caption')
        caption.alignment=WD_ALIGN_PARAGRAPH.CENTER;caption.paragraph_format.keep_with_next=False

    source=(ROOT/'docs/report_msft.md').read_text(encoding='utf-8')
    for block in re.split(r'\n\s*\n',source.strip()):
        if block.startswith('# '):doc.add_paragraph(block[2:],style='Heading 1')
        elif block.startswith('## '):doc.add_paragraph(block[3:],style='Heading 2')
        elif block.startswith('{{table:'):add_table(assets['tables'][block[8:-2]])
        elif block.startswith('{{figure:'):add_figure(block[9:-2])
        else:
            p=doc.add_paragraph(block.replace('\n',' '));p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            if re.match(r'^\d+\. ',block) and ('https://' in block or 'Методические' in block):
                p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(8);p.paragraph_format.line_spacing=1.0
                for run in p.runs:black(run,12)
    doc.core_properties.title='Прогнозирование цены акций Microsoft'
    doc.core_properties.subject='Семестровая работа по интеллектуальному анализу данных'
    doc.core_properties.author='Половинка В. А.';doc.core_properties.last_modified_by='Половинка В. А.'
    doc.core_properties.comments='';doc.core_properties.keywords='MSFT, анализ данных, прогнозирование'
    settings=doc.settings.element;update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
    args.output.parent.mkdir(parents=True,exist_ok=True);doc.save(args.output)
    assert hashlib.sha256(args.reference.read_bytes()).hexdigest()==initial_hash
    print(args.output)
    print(f'{table_no} tables, {figure_no} figures; reference unchanged')


if __name__=='__main__':main()
